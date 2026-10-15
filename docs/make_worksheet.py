"""Builds the teacher guide and student worksheet for Marathon Runner: Water Balance (model 0.5.0)."""
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
URL = "sims.mr-miner.com/play/water-balance"


def base(doc):
	st = doc.styles["Normal"]
	st.font.name = "Calibri"
	st.font.size = Pt(11)
	for s in doc.sections:
		s.top_margin = s.bottom_margin = Inches(0.7)
		s.left_margin = s.right_margin = Inches(0.8)


def shade(cell, hex_fill):
	tc = cell._tc.get_or_add_tcPr()
	sh = OxmlElement("w:shd")
	sh.set(qn("w:val"), "clear")
	sh.set(qn("w:color"), "auto")
	sh.set(qn("w:fill"), hex_fill)
	tc.append(sh)


def para(doc, text="", bold=False, italic=False, size=None, after=4):
	p = doc.add_paragraph()
	p.paragraph_format.space_after = Pt(after)
	if text:
		r = p.add_run(text)
		r.bold = bold
		r.italic = italic
		if size:
			r.font.size = Pt(size)
	return p


def rich(doc, parts, style=None, after=3):
	"""parts: list of (text, bold)."""
	p = doc.add_paragraph(style=style)
	p.paragraph_format.space_after = Pt(after)
	for text, bold in parts:
		r = p.add_run(text)
		r.bold = bold
	return p


def bullet(doc, text, bold_lead=None):
	parts = [(bold_lead, True), (text, False)] if bold_lead else [(text, False)]
	return rich(doc, parts, style="List Bullet", after=2)


def numbered(doc, text, bold_lead=None):
	parts = [(bold_lead, True), (text, False)] if bold_lead else [(text, False)]
	return rich(doc, parts, style="List Number", after=2)


def table(doc, header, rows, widths=None, row_height=None, font=10):
	# Keep the table, and the paragraph that introduces it, on one page
	if doc.paragraphs:
		doc.paragraphs[-1].paragraph_format.keep_with_next = True
	t = doc.add_table(rows=1 + len(rows), cols=len(header))
	t.style = "Table Grid"
	t.alignment = WD_TABLE_ALIGNMENT.CENTER
	for i, h in enumerate(header):
		c = t.rows[0].cells[i]
		c.text = ""
		r = c.paragraphs[0].add_run(h)
		r.bold = True
		r.font.size = Pt(font)
		shade(c, "D9E2F3")
	for ri, row in enumerate(rows, start=1):
		for ci, val in enumerate(row):
			c = t.rows[ri].cells[ci]
			c.text = ""
			r = c.paragraphs[0].add_run(val)
			r.font.size = Pt(font)
			if ci == 0:
				r.bold = True
		if row_height:
			t.rows[ri].height = Inches(row_height)
	for row in t.rows[:-1]:
		for c in row.cells:
			for cp in c.paragraphs:
				cp.paragraph_format.keep_with_next = True
	if widths:
		t.autofit = False
		for row in t.rows:
			for i, w in enumerate(widths):
				row.cells[i].width = Inches(w)
	doc.add_paragraph().paragraph_format.space_after = Pt(2)
	return t


def lines(doc, n):
	for _ in range(n):
		p = doc.add_paragraph("_" * 86)
		p.paragraph_format.space_after = Pt(6)


def heading(doc, text, level):
	h = doc.add_heading(text, level)
	h.paragraph_format.space_before = Pt(10 if level > 1 else 0)
	h.paragraph_format.space_after = Pt(4)
	return h


# =====================================================================
# TEACHER GUIDE
# =====================================================================
def teacher_guide(doc):
	heading(doc, "Marathon Runner: Water Balance — Teacher Guide", 0)
	para(doc, "A one-screen model of a runner's body during a race. Students choose how she runs and what she drinks, "
		"run the race in about a minute, and watch what happens to the water and salt in her blood.", italic=True)

	heading(doc, "At a glance", 1)
	table(doc, ["", ""], [
		["Where", URL + "   (backup: sfminer.github.io/water-balance)"],
		["Devices", "Any browser, including Chromebooks. No login, nothing saved between visits."],
		["Time", "A 4-hour race takes about 1 minute at Speed 1x (8 hours: 2 minutes). Speed 4x cuts that to 15–30 s. "
			"Plan 25–35 minutes for the worksheet."],
		["Fits", "Marathon Runner anchor phenomenon (Water Balance 5E). Best after students have seen the Medical Tent "
			"Data, before or during Explain."],
		["Pairs with", "The osmosis / cell-size simulation, which shows what happens to the cells themselves."],
	], widths=[1.2, 5.6])

	heading(doc, "What students can change", 1)
	table(doc, ["Setting", "Choices", "What it does"], [
		["Pace", "Stop / Run", "Running sweats about 8 times as fast as resting. Legs move on screen while running."],
		["How much the runner drinks", "None / Medium / Lots",
			"Medium about matches what she sweats. Lots is more than sweat plus her kidneys can get rid of."],
		["What the runner drinks", "Water / Sports drink", "A sports drink adds a little salt with the water."],
		["Race length", "4 / 6 / 8 hours", "Same rates, just longer. 4 hours is the default."],
	], widths=[1.6, 1.5, 3.7])
	bullet(doc, " Settings can be changed in the middle of a race (for example, start running, then switch to Stop halfway). "
		"Changing a setting after a race ends starts a fresh race.", "Mid-race changes.")

	heading(doc, "The buttons", 1)
	table(doc, ["Button", "What it does"], [
		["Start race / Pause / Resume", "Starts and pauses the race clock. Space does the same."],
		["Run again", "Appears when the race is over; runs the same settings from the start."],
		["Restart", "Goes back to the start of the race, keeping the settings."],
		["Speed 1x / 2x / 4x", "How fast the race clock runs. Results are the same at every speed."],
		["Keep this run", "Freezes the current sodium line on the graph as a dashed line, so the next run can be compared with it."],
		["Clear kept run", "Removes the dashed line."],
	], widths=[2.0, 4.8])

	heading(doc, "Reading the screen", 1)
	bullet(doc, " cup (drinking), sweat leaving the skin, the blood box, the ADH meter, and the kidneys with the urine arrow. "
		"Moving arrows get thicker when more is flowing.", "The body diagram (left):")
	bullet(doc, " the water level shows how much water is in her blood; the yellow diamonds show how much salt. "
		"Crowded diamonds = salty blood; spread-out diamonds = diluted blood.", "The blood box:")
	bullet(doc, " the brain's message to the kidneys. High ADH = keep water (little urine). Low ADH = let water go "
		"(more urine).", "ADH meter (purple):")
	bullet(doc, " one sentence describing what is happening right now: is the blood saltier or more diluted, "
		"is water moving into or out of the cells, and what ADH and the kidneys are doing.", "The observation line:")
	bullet(doc, " how much water and salt has gone in (Drunk) and out (Sweat, Urine) since the start. "
		"Every number is a percentage of what the blood held at the start, so 100% = as much as the blood started with.", "Water and salt so far (right):")
	bullet(doc, " salt ÷ water, turned into the mEq/L units on the Medical Tent Data. The green band is the normal "
		"range, 135–145 mEq/L.", "Blood saltiness, gauge and graph:")

	heading(doc, "When the runner collapses", 1)
	para(doc, "If a choice pushes her body too far, the race stops, her label turns red (\"Runner (collapsed)\") and the "
		"observation line says why: too much water lost, too much water held, sodium too low or sodium too high. "
		"This is intended: it keeps the screen from showing values no one could survive. Students can press Run again "
		"or change a setting. The collapse points (blood water below 65% or above 150% of its start; sodium below 110 "
		"or above 170 mEq/L) are classroom limits, not medical cut-offs.")

	heading(doc, "What to expect (answer key)", 1)
	para(doc, "Blood sodium at the end of the race, in mEq/L (normal 135–145). Small differences of ±1 can appear.")
	table(doc, ["Race", "Pace", "None", "Medium", "Lots of water", "Lots of sports drink"], [
		["4 h", "Run", "164 (high)", "139", "125 (low)", "132 (low)"],
		["4 h", "Stop", "143", "136", "113 (low)", "122 (low)"],
		["6 h", "Run", "collapses ≈5.0 h", "139", "121 (low)", "132 (low)"],
		["8 h", "Run", "collapses ≈5.0 h", "138", "117 (low)", "132 (low)"],
		["8 h", "Stop", "146", "135", "collapses ≈4.6 h", "collapses ≈6.2 h"],
	], widths=[0.6, 0.6, 1.3, 0.9, 1.6, 1.8])
	para(doc, "The big ideas students should find:", bold=True)
	numbered(doc, " Running and drinking nothing makes the blood saltier (sweat takes out water). "
		"ADH goes up and urine drops: the body saves water, but can't make up for no drinking.", "No drinking → too salty.")
	numbered(doc, " Medium drinking keeps her in or near the normal range. If she drinks a little extra, ADH falls "
		"and the kidneys get rid of it.", "About the right amount → normal.")
	numbered(doc, " Drinking lots of plain water dilutes the blood below normal, even with ADH at zero, because she "
		"drinks faster than her kidneys can get rid of water. Water then moves into the cells (including brain cells), "
		"which is why this is dangerous. This is the anchor phenomenon.", "Too much water → too diluted.")
	numbered(doc, " Resting while drinking lots is worse than running while drinking lots, because a resting "
		"runner sweats less water away.", "Stopping makes overdrinking worse.")
	numbered(doc, " A sports drink softens the drop but does not prevent it: the problem is mostly too much water, "
		"not too little salt.", "Sports drinks help a little.")

	heading(doc, "Watch for these misconceptions", 1)
	bullet(doc, " A lower sodium number means the blood is more diluted, not necessarily that she lost more salt. "
		"Point students to the Salt row: the running runner loses more salt in sweat than the resting runner, yet her "
		"blood ends up saltier, because sweat takes out even more water.", "\"Lower number = she lost more sodium.\"")
	bullet(doc, " Water in the blood and saltiness of the blood are different things. Use the blood box: same "
		"number of diamonds, more water, lower saltiness.", "Amount vs. concentration.")
	bullet(doc, " The kidneys do respond, but they have a maximum. With Lots of water, ADH reaches zero and the "
		"observation says the kidneys are at their limit.", "\"The kidneys should just fix it.\"")

	heading(doc, "Good to know", 1)
	bullet(doc, " The numbers are illustrative: they behave like the real body but are not medical measurements, "
		"and the race clock is scaled. The game does not contain the runner's own Medical Tent values; students "
		"use the printed data.", "It's a model.")
	bullet(doc, " The game pauses itself if its tab is hidden or the window is minimized. Keep the game visible "
		"while it runs.", "Keep the tab on screen.")
	bullet(doc, " If the game looks like an older version, press Ctrl+Shift+R to reload it.", "Old version showing?")
	bullet(doc, " Tab moves through the buttons (yellow outline); Space starts and pauses; Escape or tabbing past "
		"the last button leaves the game for the rest of the page.", "Keyboard.")
	bullet(doc, " Text stays readable on small screens; on a phone or a short window, scroll to see the graph.",
		"Small screens.")


# =====================================================================
# STUDENT WORKSHEET
# =====================================================================
def student_worksheet(doc):
	p = para(doc, "Name: ______________________________     Class: __________     Date: ____________", after=8)
	heading(doc, "Marathon Runner: Water Balance", 0)
	para(doc, "Why did a runner who drank plenty of water end up in the Medical Tent? In this model you control how "
		"she runs and what she drinks, then watch what happens to the water and salt in her blood.", italic=True)
	para(doc, "Open: " + URL, bold=True)

	heading(doc, "How to use the model", 1)
	numbered(doc, "Choose the Pace, How much the runner drinks, What the runner drinks and the Race length.")
	numbered(doc, "Press Start race (or Space). Use Speed 4x if you want it faster; the results are the same.")
	numbered(doc, "When the race is over, read the Blood sodium gauge and the observation sentence under the body.")
	numbered(doc, "To compare two runs, press Keep this run, change a setting, then press Run again. The kept run "
		"shows as a dashed line on the graph.")
	para(doc, "Normal blood sodium is 135–145 mEq/L (the green band). If the runner collapses, the race stops and the "
		"screen tells you why. Write that down: it is a result, not a mistake.", after=6)

	heading(doc, "Part 1: Get to know the screen", 1)
	para(doc, "Before you start a race, find each part of the body diagram. What does each one show?")
	table(doc, ["Part of the screen", "What it shows"], [
		["The cup and its arrow", ""],
		["The blood box: blue water level", ""],
		["The blood box: yellow diamonds", ""],
		["The purple ADH meter", ""],
		["The kidneys and the urine arrow", ""],
	], widths=[2.4, 4.4], row_height=0.38)

	heading(doc, "Part 2: Predict, then test", 1)
	para(doc, "Use a 4-hour race and Pace: Run unless the table says otherwise. Before each run, predict whether her "
		"blood sodium will end up too high, normal, or too low. Then run it and record what happens.")
	table(doc, ["Run", "Settings", "Prediction\n(high / normal / low)", "Sodium at end\n(mEq/L)", "ADH at end\n(high / med / low)",
		"Urine\n(a lot / some / a little)", "Collapsed? Why?"], [
		["A", "Run, None", "", "", "", "", ""],
		["B", "Run, Medium, Water", "", "", "", "", ""],
		["C", "Run, Lots, Water", "", "", "", "", ""],
		["D", "Run, Lots, Sports drink", "", "", "", "", ""],
		["E", "Stop, Lots, Water", "", "", "", "", ""],
		["F", "8 hours, Run, Lots, Water", "", "", "", "", ""],
	], widths=[0.4, 1.5, 1.0, 0.9, 1.0, 1.0, 1.1], row_height=0.42, font=9)
	para(doc, "1. Which runs ended outside the normal range? Which direction was each one?")
	lines(doc, 2)
	para(doc, "2. In Run C, what did the observation sentence say about ADH and the kidneys? Why couldn't the kidneys "
		"bring her sodium back to normal?")
	lines(doc, 2)

	heading(doc, "Part 3: Compare two runners", 1)
	para(doc, "Run C again and press Keep this run. Then change Pace to Stop (Run E) and press Run again.")
	para(doc, "3. Which runner's blood was more diluted at the end, the one running or the one stopped? "
		"Use what sweat does to explain why.")
	lines(doc, 2)
	para(doc, "4. Run A (Run, None) and then Stop with None. Copy the Salt row of the \"Water and salt so far\" table.")
	table(doc, ["", "Salt lost in sweat", "Sodium at end (mEq/L)"], [
		["Run, None", "", ""],
		["Stop, None", "", ""],
	], widths=[1.6, 2.2, 2.2], row_height=0.35)
	para(doc, "Which runner lost more salt? Which one has saltier blood? How can both be true? "
		"(Hint: look at the Water row too.)")
	lines(doc, 2)
	para(doc, "5. Compare Run C (water) and Run D (sports drink). Did the sports drink prevent the problem? "
		"What does that tell you about the main cause?")
	lines(doc, 2)

	heading(doc, "Part 4: Back to the Medical Tent", 1)
	para(doc, "Use your printed Medical Tent Data.")
	table(doc, ["", "Blood sodium (mEq/L)", "Normal range?"], [
		["Our runner (from the Medical Tent Data)", "", ""],
		["The model's closest run (which one?)", "", ""],
	], widths=[3.0, 1.9, 1.9], row_height=0.38)
	para(doc, "6. Which settings in the model come closest to what happened to our runner? Describe how she probably "
		"ran and drank.")
	lines(doc, 2)
	para(doc, "7. When her blood became diluted, which way did water move: into her cells or out of them? "
		"(Check the observation sentence.) Why would that be dangerous for brain cells?")
	lines(doc, 2)

	heading(doc, "Part 5: Explain it", 1)
	para(doc, "8. Claim, evidence, reasoning: Why did a runner who drank lots of water end up with low blood sodium? "
		"Use at least two numbers from the model and the words ADH, kidneys and diluted.")
	lines(doc, 5)
	para(doc, "9. What advice would you give a runner about drinking during a long race? What does this model leave "
		"out that a real runner might also need to think about?")
	lines(doc, 3)
	para(doc, "Remember: this is a simplified model. Its numbers act like a real body, but they are not medical "
		"measurements.", italic=True, size=9)


which = sys.argv[2]
d = Document()
base(d)
(teacher_guide if which == "teacher" else student_worksheet)(d)
d.save(OUT)
print("wrote", OUT)
